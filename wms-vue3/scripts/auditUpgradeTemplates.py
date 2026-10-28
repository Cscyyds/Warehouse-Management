"""Verify versioned blank templates against the local backend Excel parser."""
from __future__ import annotations

import ast
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT.parent / 'nuomi_wms'
sys.path.insert(0, str(BACKEND))

from openpyxl import load_workbook  # backend dependency
from app.utils.excel_import import parse_excel_file


def constant_lists(source: Path) -> dict[str, list[str]]:
    result = {}
    for node in ast.parse(source.read_text(encoding='utf-8')).body:
        target = node.target if isinstance(node, ast.AnnAssign) else (
            node.targets[0] if isinstance(node, ast.Assign) and len(node.targets) == 1 else None)
        if isinstance(target, ast.Name):
            try:
                value = ast.literal_eval(node.value)
            except (ValueError, TypeError):
                continue
            if isinstance(value, list) and all(isinstance(x, str) for x in value):
                result[target.id] = value
    return result


def check(path: Path, expected: list[str], text_columns: list[str], last_row: int) -> None:
    headers, rows, count = parse_excel_file(file_bytes=path.read_bytes(), required_headers=expected)
    if headers != expected or rows or count:
        raise AssertionError(f'{path.name}: 标题或空模板内容不符合预期')
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        for name in text_columns:
            column = expected.index(name) + 1
            for row in (2, last_row):
                if sheet.cell(row, column).number_format != '@':
                    raise AssertionError(f'{path.name}: {name} 第 {row} 行不是文本格式')
    finally:
        book.close()
    print(f'{path.name}: {len(headers)} 列，空数据及编码格式通过')


def main() -> None:
    product = constant_lists(BACKEND / 'app/api/v1/endpoints/tenant_product_management.py')
    stock = constant_lists(BACKEND / 'app/api/v1/endpoints/platform_wms_stock_management.py')
    tenant_dir = ROOT / 'wms-vue3/public/templates'
    admin_dir = ROOT / 'admin-web/public/templates'
    original = load_workbook(tenant_dir / 'product-import-template.xlsx', read_only=True, data_only=True)
    try:
        original_headers = list(next(original.active.iter_rows(min_row=1, max_row=1, values_only=True)))
    finally:
        original.close()

    version = '-template-v20261027.xlsx'
    check(tenant_dir / ('product-import' + version), original_headers + ['客户类型名称', '客户类型销售价'], ['品号', '供应商型号'], 5001)
    check(tenant_dir / ('product-update' + version), product['_PRODUCT_EXPORT_HEADERS'][:25], ['产品编码', '品号'], 5001)
    supplier = product['_PRODUCT_SUPPLIER_REQUIRED_HEADERS'] + ['供应商型号', '供应商预设采购价']
    sale_price = product['_PRODUCT_SALE_PRICE_REQUIRED_HEADERS'] + ['销售价格', '备注']
    check(tenant_dir / ('product-supplier' + version), supplier, ['产品编码', '供应商型号'], 5001)
    check(tenant_dir / ('product-sale-price' + version), sale_price, ['产品编码'], 5001)
    check(admin_dir / ('product-barcode-stock' + version), stock['_BARCODE_STOCK_REQUIRED_HEADERS'], ['产品编码', '位置编码'], 2001)
    check(admin_dir / ('merge-package-stock' + version), stock['_MERGE_STOCK_REQUIRED_HEADERS'] + ['合包长度', '长度单位'], ['产品编码', '位置编码'], 2001)
    check(admin_dir / ('plastic-box-stock' + version), stock['_BOX_STOCK_REQUIRED_HEADERS'], ['产品编码', '塑料盒编码'], 2001)


if __name__ == '__main__':
    main()
